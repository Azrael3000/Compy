
#  ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
#
#           ━━━━━━━━━━━━━
#            ┏┓┏┓┳┳┓┏┓┓┏
#            ┃ ┃┃┃┃┃┃┃┗┫
#            ┗┛┗┛┛ ┗┣┛┗┛
#           ━━━━━━━━━━━━━
#
#  ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
#
#  Competition organization tool
#  for freediving competitions.
#
#  Copyright 2023 - Arno Mayrhofer
#
#  Licensed under the GNU AGPL
#
#  ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
#
#  Authors:
#
#  - Arno Mayrhofer
#
#  ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

try:
    import pandas as pd
except ImportError:
    print("Could not find pandas. Install with 'pip3 install pandas'")
    exit(-1)
import logging
try:
    import qrcode
except ImportError:
    print("Could not find qrcode. Install with 'pip3 install qrcode'")
    exit(-1)
import hashlib
import math
import random
import os
import sqlite3
from datetime import datetime, timedelta, time, timezone
from io import BytesIO
import base64
from PIL import Image
import regex
import sys

try:
    import country_converter
except ImportError:
    print("Could not find country_converter. Install with 'pip3 install country_converter'")
    exit(-1)

import athlete
from compy_aida_api import AidaApiClient, AidaApiError
from compy_config import CompyConfig
from compy_constants import (INVALID_DATE, INVALID_TIME, POOL_DISCIPLINES,
                             DEPTH_DISCIPLINES, DISCIPLINES, FEDERATIONS)
from compy_pdf import PdfReportMixin

import compy_utilities as u
from openpyxl import load_workbook

class CompyData(PdfReportMixin):

    version_ = None

    def __init__(self, db, app, comp_id = None, published_only = False):
        """Per-request competition data object.

        One CompyData is created for every request that needs competition
        state (see CompyFlask.getData). It must never be stored on the
        application object or shared between requests - a shared instance
        would make "the currently open competition" global to the server,
        breaking concurrent admin/judge/clock/results pages.

        comp_id None creates an empty (unloaded) object, an integer loads
        that competition. With published_only=True only competitions whose
        results are published can be loaded (for the public results pages).
        """
        self.id_ = None
        self.db_ = db
        self.app_ = app
        self.name_ = "undefined"
        self.special_ranking_name_ = "Newcomer"
        self.config_ = CompyConfig()
        self.lane_style_ = "numeric"
        self.comp_type_ = FEDERATIONS[0]
        self.comp_file_ = ''
        self.start_date_ = None
        self.end_date_ = None
        self.nrs_ = None
        self.tiered_records_ = None
        self.sponsor_img_ = None
        self.disciplines_ = 0
        self.selected_country_ = None
        self.publish_results_ = False
        self.aida_event_id_ = None
        self.aida_api_key_ = None

        if comp_id is not None:
            try:
                comp_id = int(comp_id)
            except (TypeError, ValueError):
                comp_id = None
            if comp_id is not None and published_only:
                comp_id, _ = self.cleanCompIdPublished(comp_id)
            if comp_id is not None:
                self.load(comp_id)

    @classmethod
    def ensureDefaultCompetition(cls, db, app):
        """Create the default competition on a fresh database.

        The admin frontend loads competition 1 on startup, so at least one
        competition has to exist. Called once at server start (inside an app
        context); requests never create competitions implicitly.
        """
        data = cls(db, app)
        if db.execute("SELECT id FROM competition") is None:
            data.save()
        return data

    @property
    def version(self):
        if self.version_ is None:
            base_path = os.path.dirname(os.path.realpath(__file__))
            with open(os.path.join(base_path, 'VERSION'), 'r') as f:
                self.version_ = str(f.read().strip())
        return self.version_

    @property
    def name(self):
        return self.name_

    @property
    def isValid(self):
        return self.id_ is not None

    @property
    def special_ranking_name(self):
        return self.special_ranking_name_

    @property
    def publish_results(self):
        return self.publish_results_

    @property
    def config(self):
        return self.config_

    @property
    def lane_style(self):
        return self.lane_style_

    def changeLaneStyle(self, lane_style):
        allowed_lane_styles = ["numeric", "alphabetic"]
        if lane_style not in allowed_lane_styles:
            return 1
        self.lane_style_ = lane_style
        self.save()
        return 0

    @property
    def comp_type(self):
        return self.comp_type_

    def changeCompType(self, comp_type):
        if comp_type not in FEDERATIONS:
            return 1
        self.comp_type_ = comp_type
        # the record caches are scoped to the federation
        self.nrs_ = None
        self.tiered_records_ = None
        self.save()
        return 0

    def laneStyleConverter(self, lane, invert=False):
        if invert:
            if self.lane_style_ == "alphabetic":
                return ord(lane) - 64
            else: # numeric
                return int(lane)
        else:
            if self.lane_style_ == "alphabetic":
                return chr(lane + 64)
            else: # numeric
                return str(lane)

    @property
    def comp_file(self):
        return self.comp_file_

    @property
    def number_of_athletes(self):
        db_out = self.db_.execute("SELECT COUNT(*) FROM competition_athlete WHERE competition_id=?", self.id_)
        if db_out is None:
            return 0
        else:
            return db_out[0][0]

    @property
    def start_date(self):
        return self.start_date_

    @property
    def end_date(self):
        return self.end_date_

    @property
    def disciplines(self):
        return self.getDisciplinesFromInt(self.disciplines_)

    @property
    def sponsor_img_data(self):
        if self.sponsor_img_ is None:
            return ""
        else:
            return self.sponsor_img_["data"]

    @property
    def countries(self):
        if self.id_ is None:
            return None
        c_data = self.db_.execute('''SELECT DISTINCT athlete.country FROM athlete
                                     INNER JOIN competition_athlete
                                     ON athlete.id==competition_athlete.athlete_id
                                     WHERE competition_athlete.competition_id==?''',
                         self.id_)
        if c_data is None:
            return []
        return [c[0] for c in c_data]

    @property
    def sponsor_img_width(self):
        if self.sponsor_img_ is None:
            return 0
        else:
            aspect_ratio = self.sponsor_img_["aspect_ratio"]
            if aspect_ratio < 1:
                return 19.*self.sponsor_img_["aspect_ratio"]
            else:
                return 19.

    @property
    def sponsor_img_height(self):
        if self.sponsor_img_ is None:
            return 0
        else:
            aspect_ratio = self.sponsor_img_["aspect_ratio"]
            if aspect_ratio > 1:
                return 5./self.sponsor_img_["aspect_ratio"]
            else:
                return 5.

    @property
    def selected_country(self):
        if self.selected_country_ is None:
            return "none"
        else:
            return self.selected_country_

    @property
    def nr(self):
        if self.nrs_ is None:
            nrs = self.db_.execute('''SELECT country, class, gender, discipline, value
                                      FROM records WHERE federation=? AND tier='NR' ''',
                                   self.comp_type)
            if nrs is None:
                return None
            self.nrs_ = {}
            for nr in nrs:
                self.nrs_[u.NR(self.comp_type, nr[0], nr[1], nr[2], nr[3])] = nr[4]
        return self.nrs_

    def changeSponsorImage(self, img_content):
        self.sponsor_img_ = {}
        img_base64 = base64.b64encode(img_content).decode('utf-8')
        self.sponsor_img_["data"] = 'data:image/png;base64,' + img_base64
        img = Image.open(BytesIO(img_content))
        width, height = img.size # in pixels
        self.sponsor_img_["aspect_ratio"] = float(width)*5./float(height)/19. # < 1 if too high, > 1 if too wide
        self.save()

    def compFileChange(self, comp_file):
        self.comp_file_ = comp_file
        self.refresh()

    def refresh(self):
        if not os.path.exists(self.comp_file_):
            logging.error("File '%s' does not exist", self.comp_file_)
            return
        # run the whole re-import atomically so that concurrent readers never
        # see a competition with deleted athletes but no new data yet
        with self.db_.transaction():
            self.refreshImpl()

    def refreshImpl(self):
        # special ranking ids of athletes
        sr_ids = None
        if self.id_ is not None:
            sr_ids = self.db_.execute(
                '''SELECT a.id FROM athlete a
                   INNER JOIN competition_athlete ca ON a.id == ca.athlete_id
                   WHERE ca.special_ranking AND ca.competition_id==?''',
                self.id_)
            self.db_.execute('''DELETE FROM start
                                WHERE competition_athlete_id IN (
                                    SELECT competition_athlete_id FROM start
                                    INNER JOIN competition_athlete
                                    ON start.competition_athlete_id == competition_athlete.id
                                    WHERE competition_athlete.competition_id == ?)''',
                             self.id_)
            self.db_.execute("DELETE FROM competition_athlete WHERE competition_id=?", self.id_)
        # read first sheet (start & end date)
        df = pd.read_excel(self.comp_file_, sheet_name="Event")
        i = 0
        for l in df[df.keys()[0]]:
            if l == "Starts:":
                self.start_date_ = df[df.keys()[1]][i]
            if l == "Ends:":
                self.end_date_ = df[df.keys()[1]][i]
            if l == "Disciplines:":
                self.disciplines_ = self.disciplineListToInt(df[df.keys()[1]][i].split(","))
            i += 1
        logging.debug("Start date: %s. End date: %s", self.start_date_, self.end_date_)

        # read second sheet (list of athletes)
        df = pd.read_excel(self.comp_file_, sheet_name="Athletes and Judges", skiprows=1)
        for i,r in df.iterrows():
            a = athlete.Athlete.fromArgs(r['Id'], r['FirstName'], r['LastName'], r['Gender'], r['Country'], r['Club'] if 'Club' in r else "", self.db_)
            logging.debug("Athlete: %s %s %s %s %s", r['Id'], r['FirstName'], r['LastName'], r['Gender'], r['Country'])
            a.associateWithComp(self.id_)
        logging.debug("Number of athletes: %d", self.number_of_athletes)

        if sr_ids is not None:
            for sr in sr_ids:
                self.setRegistration(sr[0], True, "specialranking", warn=False)

        self.save()

        ap_lambda = lambda x, y, d, self: None if math.isnan(x) else (int(x)*60+float(y) if d=="STA" else float(x))
        self.db_.execute('''DELETE FROM block WHERE competition_id == ?''',
                         self.id_)
        for day in self.getDaysExcel(self.start_date, self.end_date):
            df = pd.read_excel(self.comp_file_, sheet_name=day, skiprows=1)
            blocks = {}
            day_db = day.replace('-', '')
            for i,r in df.iterrows():
                dis = r['Discipline']
                if not dis in blocks.keys():
                    blocks[dis] = self.db_.insert('''INSERT INTO block
                                        (competition_id, day, disciplines)
                                        VALUES (?, ?, ?)''',
                                     (self.id_, day_db, self.disciplineListToInt([dis])))
            for i,r in df.iterrows():
                aida_id = r['Diver Id']
                ca_id = self.db_.execute('''SELECT competition_athlete.id FROM competition_athlete
                                            INNER JOIN athlete
                                            ON competition_athlete.athlete_id == athlete.id
                                            WHERE athlete.aida_id=? AND competition_athlete.competition_id == ?''',
                                         (aida_id, self.id_))
                dis = r['Discipline']
                ap = ap_lambda(r['Meters or Min'], r['Sec(STA only)'], dis, self)
                ot = self.parseTime(r['OT'])
                lane = int(r['Zone'])
                rp = ap_lambda(r['Meters or Min.1'], r['Sec(STA only).1'], dis, self)
                card = r['Card']
                pen_other = float(r['Pen(other)']) if not math.isnan(r['Pen(other)']) else 0.
                penalty = float(r['Pen(UNDER AP)']) + pen_other
                remarks = r['Remarks']
                if remarks == "DNS":
                    rp = float('nan')
                    card = "nan"
                    penalty = "nan"
                block = blocks[dis]
                if rp is not None:
                    self.db_.execute('''INSERT INTO start
                                        (competition_athlete_id, discipline, lane, OT, AP,
                                         rp, card, penalty, remarks, block)
                                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                                     (ca_id[0][0], dis, lane, u.convTime(ot), ap, rp, card, penalty, remarks, block))
                else:
                    self.db_.execute('''INSERT INTO start
                                        (competition_athlete_id, discipline, lane, OT, AP, block)
                                        VALUES (?, ?, ?, ?, ?, ?)''',
                                     (ca_id[0][0], dis, lane, u.convTime(ot), ap, block))


    def setRegistration(self, athlete_id, is_checked, change_type, warn=True):
        found = False
        if self.number_of_athletes == 0:
            logging.warning("Data not initialized yet in setRegistration")
            return 1
        try:
            type_map = {'eligiblenational': 'eligible_national', 'specialranking': 'special_ranking', 'paid': 'paid', 'medicalchecked': 'medical_checked', 'registered': 'registered'}
            if not change_type in type_map:
                return 0
            self.db_.execute(
                '''UPDATE competition_athlete SET {}=?
                   WHERE competition_id==? AND athlete_id==?'''.format(type_map[change_type]),
                (is_checked, self.id_, athlete_id))
            return 0
        except sqlite3.Error as e:
            if warn:
                logging.warning("Tried setting " + change_type + " (" + str(is_checked) + ") to athlete with id '" + athlete_id + "' but this id could not be found")
            return 1

    #
    # ---- AIDA International API integration (read only) ----
    #
    # The API key belongs to the competition (AIDA generates a new key per
    # event), so both settings live in the competition table and travel with
    # the saved competition. All network I/O happens before any database
    # write and the whole sync is applied in a single transaction, so a
    # failed or malformed fetch never leaves partial state behind (there is
    # no AIDA test environment - the excel workflow stays available as the
    # fallback at all times). The sync never deletes local data and never
    # touches locally entered results or OTs.
    #

    @property
    def aida_event_id(self):
        return self.aida_event_id_

    @property
    def has_aida_api_key(self):
        return bool(self.aida_api_key_)

    def getAidaStatus(self, data):
        data["aida_event_id"] = self.aida_event_id_
        data["aida_has_key"] = self.has_aida_api_key

    def setAidaSettings(self, event_id, api_key):
        """Store the AIDA event id and API key for this competition.

        An empty api_key keeps the already stored key, so the event id can
        be changed without re-entering the key (the frontend never learns
        the stored key, it only knows whether one exists).
        """
        if self.id_ is None:
            return 1, {"error_msg": "No competition loaded"}
        if event_id is None or str(event_id).strip() == "":
            self.aida_event_id_ = None
        else:
            try:
                self.aida_event_id_ = int(str(event_id).strip())
            except (TypeError, ValueError):
                return 1, {"error_msg": "The AIDA event id must be a number"}
        if api_key is not None and str(api_key).strip() != "":
            self.aida_api_key_ = str(api_key).strip()
        self.save()
        data = {}
        self.getAidaStatus(data)
        return 0, data

    def aidaClient(self):
        return AidaApiClient(self.aida_api_key_)

    def checkAidaConfigured(self):
        if self.id_ is None:
            return {"error_msg": "No competition loaded"}
        if self.aida_event_id_ is None or not self.has_aida_api_key:
            return {"error_msg": "Set the AIDA event id and API key first"}
        return None

    def testAidaConnection(self, client=None):
        """Validate key, scope, event binding and connectivity in one call.

        Uses the cheapest read endpoint (list of days) and returns the event
        name and days so the admin can verify the key belongs to the right
        event days before the competition.
        """
        error = self.checkAidaConfigured()
        if error is not None:
            return 1, error
        try:
            client = client if client is not None else self.aidaClient()
            payload = client.getDays(self.aida_event_id_)
        except AidaApiError as e:
            return 1, {"error_msg": str(e)}
        days = [str(d.get("date"))[:10] for d in payload["days"]]
        return 0, {"event_name": payload.get("eventName", "unknown"),
                   "days": days}

    def syncFromAida(self, client=None):
        """Read-only sync of this competition from the AIDA API.

        Pulls days, registrations and start lists, then applies everything
        as one atomic upsert. Athletes are matched by their AIDA UUID,
        starts by their AIDA start id (falling back to athlete+discipline+
        day). Local-only athletes and starts are reported but never
        deleted; locally entered OTs, PBs, clubs and results are preserved.
        Returns (0, report) on success and (1, {"error_msg": ...}) if
        anything failed - in the failure case the database is unchanged.
        """
        error = self.checkAidaConfigured()
        if error is not None:
            return 1, error

        # 1) fetch and validate everything before touching the database
        try:
            client = client if client is not None else self.aidaClient()
            days_payload = client.getDays(self.aida_event_id_)
            prereg_payload = client.getPreregistrations(self.aida_event_id_)
            start_lists = []
            for day in days_payload["days"]:
                if day.get("isRestDay"):
                    continue
                start_lists.append(
                    (day, client.getStartList(self.aida_event_id_, day["id"])))
        except AidaApiError as e:
            return 1, {"error_msg": str(e)}

        report = {"athletes_added": 0, "athletes_updated": 0,
                  "starts_added": 0, "starts_updated": 0,
                  "days_synced": 0, "pbs_filled": 0,
                  "warnings": [], "only_local": []}
        athletes_in = self.normalizeAidaAthletes(
            prereg_payload["preregistrations"], report["warnings"])
        if len(athletes_in) == 0:
            return 1, {"error_msg": "AIDA returned no valid registrations; "
                                    "nothing was changed"}
        pbs = self.fetchAidaPbs(client, athletes_in, report["warnings"])

        # 2) apply atomically; any error rolls the whole sync back
        try:
            with self.db_.transaction():
                self.applyAidaSync(days_payload, athletes_in, start_lists,
                                   report)
                self.applyAidaPbs(pbs, report)
        except sqlite3.Error as e:
            logging.error("AIDA sync failed, rolled back: %s", e)
            return 1, {"error_msg": "Database error during sync; "
                                    "no changes were applied"}
        return 0, report

    def normalizeAidaAthletes(self, entries, warnings):
        """Validate the registration list into Compy's athlete fields.

        Invalid entries are skipped with a warning instead of failing the
        whole sync. Entries without an athleteId cannot be matched reliably
        and are skipped as well.
        """
        athletes_in = []
        for entry in entries:
            if not isinstance(entry, dict):
                continue
            aida_id = str(entry.get("athleteId") or "").strip()
            first_name = str(entry.get("firstName") or "").strip()
            last_name = str(entry.get("lastName") or "").strip()
            label = (first_name + " " + last_name).strip() or aida_id or "?"
            if aida_id == "" or first_name == "" or last_name == "":
                warnings.append("Skipped registration with incomplete "
                                "identity: " + label)
                continue
            gender = str(entry.get("diverGender") or "").strip().upper()[:1]
            if gender not in ("M", "F"):
                warnings.append("Skipped " + label + ": unknown gender")
                continue
            country = self.aidaCountryToIoc(entry)
            if country is None:
                warnings.append("Skipped " + label + ": unknown nationality")
                continue
            status = entry.get("status") or {}
            documents = entry.get("documents") or {}
            athletes_in.append({
                "aida_id": aida_id,
                "first_name": first_name,
                "last_name": last_name,
                "gender": gender,
                "country": country,
                "registered": bool(status.get("code") == "APP"),
                "paid": bool(documents.get("paymentReceipt")),
                "medical_checked": bool(documents.get("medicalCertificate")),
            })
        return athletes_in

    # class level cache: building a CountryConverter is expensive
    country_converter_ = None

    @classmethod
    def countryConverter(cls):
        if cls.country_converter_ is None:
            cls.country_converter_ = country_converter.CountryConverter()
        return cls.country_converter_

    def aidaCountryToIoc(self, entry):
        """AIDA sends ISO alpha-3 / country names; Compy stores IOC codes."""
        cc = self.countryConverter()
        for value in (entry.get("diverNationalityAbrvAlpha3"),
                      entry.get("diverNationality")):
            if not value:
                continue
            ioc = cc.convert(names=str(value), to="IOC", not_found=None)
            if isinstance(ioc, str) and self.cleanCountry(ioc) is not None:
                return ioc
        # last resort: the alpha-3 code itself if it looks like an IOC code
        return self.cleanCountry(entry.get("diverNationalityAbrvAlpha3"))

    def applyAidaSync(self, days_payload, athletes_in, start_lists, report):
        # -- athletes: upsert by AIDA UUID, preserve locally stored club --
        athlete_ids = {}
        for a in athletes_in:
            existing = self.db_.execute(
                "SELECT id, club FROM athlete WHERE aida_id=?", a["aida_id"])
            club = existing[0][1] if existing is not None else ""
            in_comp = None if existing is None else self.db_.execute(
                '''SELECT id FROM competition_athlete
                   WHERE athlete_id=? AND competition_id=?''',
                (existing[0][0], self.id_))
            ath = athlete.Athlete.fromArgs(a["aida_id"], a["first_name"],
                                           a["last_name"], a["gender"],
                                           a["country"], club, self.db_)
            ath.associateWithComp(self.id_)
            athlete_ids[a["aida_id"]] = ath.id
            report["athletes_added" if in_comp is None
                   else "athletes_updated"] += 1
            self.db_.execute(
                '''UPDATE competition_athlete
                   SET registered=?, paid=?, medical_checked=?
                   WHERE competition_id=? AND athlete_id=?''',
                (a["registered"], a["paid"], a["medical_checked"],
                 self.id_, ath.id))

        # -- athletes that exist locally but not on AIDA: report, never delete --
        api_ids = set(athlete_ids.keys())
        local = self.db_.execute(
            '''SELECT a.aida_id, a.first_name, a.last_name FROM athlete a
               INNER JOIN competition_athlete ca ON ca.athlete_id == a.id
               WHERE ca.competition_id == ?''', self.id_)
        if local is not None:
            report["only_local"] = [l[1] + " " + l[2] for l in local
                                    if l[0] not in api_ids]

        # -- competition dates from the day list --
        dates = sorted(set(self.cleanDay(str(d.get("date"))[:10], True)
                           for d in days_payload["days"]) - {INVALID_DATE})
        if len(dates) > 0:
            self.start_date_ = dates[0]
            self.end_date_ = dates[-1]

        # -- blocks and starts per day --
        for day, payload in start_lists:
            day_date = self.cleanDay(str(day.get("date"))[:10], True)
            if day_date == INVALID_DATE:
                report["warnings"].append("Skipped day with invalid date: "
                                          + str(day.get("date")))
                continue
            report["days_synced"] += 1
            day_db = day_date.replace('-', '')
            blocks = {}
            for entry in payload["startList"]:
                self.applyAidaStart(entry, day, day_db, blocks, report)
        self.save()

    def applyAidaStart(self, entry, day, day_db, blocks, report):
        if not isinstance(entry, dict):
            return
        dis = str(entry.get("disciplineAbrv") or "").strip().upper()
        label = (str(entry.get("diverFirstName") or "") + " "
                 + str(entry.get("diverLastName") or "")).strip()
        if dis not in DISCIPLINES:
            report["warnings"].append("Skipped start of " + label
                                      + ": unknown discipline '" + dis + "'")
            return
        self.disciplines_ |= self.disciplineListToInt([dis])

        # block per (day, discipline), like the excel import; keep existing
        # blocks so start list edits survive a re-sync
        if dis not in blocks:
            dis_int = self.disciplineListToInt([dis])
            existing_block = self.db_.execute(
                '''SELECT id FROM block
                   WHERE competition_id=? AND day=? AND disciplines=?''',
                (self.id_, day_db, dis_int))
            if existing_block is not None:
                blocks[dis] = existing_block[0][0]
                self.db_.execute("UPDATE block SET aida_day_id=? WHERE id=?",
                                 (day.get("id"), blocks[dis]))
            else:
                blocks[dis] = self.db_.insert(
                    '''INSERT INTO block
                       (competition_id, day, disciplines, aida_day_id)
                       VALUES (?, ?, ?, ?)''',
                    (self.id_, day_db, dis_int, day.get("id")))

        # the start list carries no athlete UUID, so match by name+gender
        gender = str(entry.get("diverGender") or "").strip().upper()[:1]
        ca_id = self.db_.execute(
            '''SELECT ca.id FROM competition_athlete ca
               INNER JOIN athlete a ON ca.athlete_id == a.id
               WHERE a.first_name=? AND a.last_name=? AND a.gender=?
               AND ca.competition_id=?''',
            (str(entry.get("diverFirstName") or "").strip(),
             str(entry.get("diverLastName") or "").strip(),
             gender, self.id_))
        if ca_id is None:
            report["warnings"].append("Skipped start of " + label
                                      + ": not in the registration list")
            return
        ca_id = ca_id[0][0]

        ap = self.aidaPerformance(entry.get("announcedPerformance"),
                                  entry.get("performanceMeasurement"), dis)
        lane = None
        try:
            if entry.get("lane") is not None:
                lane = max(1, int(entry["lane"]))
        except (TypeError, ValueError):
            lane = None
        aida_start_id = entry.get("startId")

        # match an existing start by AIDA start id first, then by
        # athlete+discipline+day (first sync of an excel-seeded comp)
        existing = None
        if aida_start_id is not None:
            existing = self.db_.execute(
                '''SELECT s.id FROM start s
                   INNER JOIN competition_athlete ca
                   ON ca.id == s.competition_athlete_id
                   WHERE s.aida_start_id=? AND ca.competition_id=?''',
                (aida_start_id, self.id_))
        if existing is None:
            existing = self.db_.execute(
                '''SELECT s.id FROM start s
                   INNER JOIN block b ON b.id == s.block
                   WHERE s.competition_athlete_id=? AND s.discipline=?
                   AND b.day=?''',
                (ca_id, dis, day_db))
        if existing is not None:
            # update announcement data only; OT, PB and results stay local
            self.db_.execute(
                '''UPDATE start SET AP=?, aida_start_id=?,
                   lane=COALESCE(?, lane) WHERE id=?''',
                (ap, aida_start_id, lane, existing[0][0]))
            report["starts_updated"] += 1
        else:
            self.db_.execute(
                '''INSERT INTO start
                   (competition_athlete_id, discipline, lane, block, OT, AP,
                    aida_start_id)
                   VALUES (?, ?, ?, ?, ?, ?, ?)''',
                (ca_id, dis, lane if lane is not None else 1, blocks[dis],
                 u.convTime(INVALID_TIME), ap, aida_start_id))
            report["starts_added"] += 1

    def fetchAidaPbs(self, client, athletes_in, warnings):
        """Fetch personal bests per athlete before the sync transaction.

        A failed profile fetch only skips that athlete (with a warning), the
        sync itself proceeds. Athletes that are already in the competition
        with no empty PB left are skipped entirely, so repeated syncs do not
        query one profile per athlete again. Returns
        {aida_id: {discipline: pb}}.
        """
        pbs = {}
        consecutive_failures = 0
        for a in athletes_in:
            if not self.needsAidaPbFetch(a["aida_id"]):
                continue
            try:
                payload = client.getAthleteProfile(self.aida_event_id_,
                                                   a["aida_id"])
            except AidaApiError as e:
                consecutive_failures += 1
                warnings.append("No AIDA profile for " + a["first_name"] + " "
                                + a["last_name"] + ": " + str(e))
                if consecutive_failures >= 3:
                    # the endpoint is probably unavailable for this event
                    # (as of 2026-08 it 404s for every athlete while the
                    # event has no start lists): collapse the per-athlete
                    # warnings into one and stop probing
                    del warnings[-consecutive_failures:]
                    warnings.append(
                        "Could not fetch athlete profiles from AIDA ("
                        + str(consecutive_failures) + " attempts in a row "
                        "failed, last error: " + str(e) + "); personal "
                        "bests were not pre-filled")
                    break
                continue
            consecutive_failures = 0
            per_dis = {}
            for entry in payload["profile"]["personalBests"]:
                if not isinstance(entry, dict):
                    continue
                dis = str(entry.get("disciplineAbrv") or "").strip().upper()
                pb = self.aidaPerformance(entry.get("reportedPerformance"),
                                          entry.get("performanceMeasurement"),
                                          dis)
                if dis in DISCIPLINES and pb is not None:
                    per_dis[dis] = pb
            if len(per_dis) > 0:
                pbs[a["aida_id"]] = per_dis
        return pbs

    def needsAidaPbFetch(self, aida_id):
        """A profile is only worth fetching for athletes that are new to
        the competition or still have a start without a PB. A start added
        by the current sync gets its PB on the next sync."""
        in_comp = self.db_.execute(
            '''SELECT ca.id FROM competition_athlete ca
               INNER JOIN athlete a ON a.id == ca.athlete_id
               WHERE a.aida_id=? AND ca.competition_id=?''',
            (aida_id, self.id_))
        if in_comp is None:
            return True
        empty = self.db_.execute(
            '''SELECT s.id FROM start s WHERE s.competition_athlete_id=?
               AND (s.PB IS NULL OR s.PB='')''', in_comp[0][0])
        return empty is not None

    def applyAidaPbs(self, pbs, report):
        """Fill start PBs that are still empty; local PBs are never touched."""
        for aida_id, per_dis in pbs.items():
            for dis, pb in per_dis.items():
                empty = self.db_.execute(
                    '''SELECT s.id FROM start s
                       INNER JOIN competition_athlete ca
                       ON ca.id == s.competition_athlete_id
                       INNER JOIN athlete a ON a.id == ca.athlete_id
                       WHERE a.aida_id=? AND ca.competition_id=?
                       AND s.discipline=? AND (s.PB IS NULL OR s.PB='')''',
                    (aida_id, self.id_, dis))
                if empty is None:
                    continue
                for row in empty:
                    self.db_.execute("UPDATE start SET PB=? WHERE id=?",
                                     (pb, row[0]))
                    report["pbs_filled"] += 1

    def aidaPerformance(self, value, measurement, discipline):
        """AIDA performance to Compy units (m, or s for STA).

        STA values arrive either as a number or, from the records
        endpoint, as "m:ss" (the live API sends e.g. "10:12" with
        performanceMeasurement "time").
        """
        if value is None or value == "":
            return None
        if ":" in str(value):
            parts = str(value).split(":")
            try:
                return float(int(parts[0])) * 60. + float(int(parts[1]))
            except (IndexError, ValueError):
                return None
        try:
            perf = float(value)
        except (TypeError, ValueError):
            return None
        if discipline == "STA" and "min" in str(measurement or "").lower():
            perf *= 60.
        return perf

    def getSavedCompetitions(self):
        saved_comp_info = []
        comps = self.db_.execute("SELECT id, name, save_date FROM competition")
        if comps is None:
            return None
        for comp in comps:
            saved_comp_info.append({"comp_id": comp[0], "name": comp[1], "save_date": comp[2]})
        return sorted(saved_comp_info, key=lambda ci: ci["save_date"], reverse=True)

    def changeName(self, new_name, overwrite):
        comp_id = self.db_.execute("SELECT id FROM competition WHERE name=?", new_name)
        file_exists = False
        comps = []
        if not comp_id is None and not overwrite:
            file_exists = True
        else:
            self.name_ = new_name
            if comp_id is None:
                self.id_ = None
            else:
                self.id_ = comp_id[0][0]
            self.save()
            comps = self.getSavedCompetitions()
        return {'file_exists': file_exists, 'name': self.name_, 'comp_id': self.id_, 'competitions': comps}

    def save(self):
        # force version update
        self.version_ = None
        if self.id_ is None:
            self.id_ = self.db_.insert('''INSERT INTO competition
                                (name, save_date, version, lane_style, comp_type, comp_file, start_date,
                                 end_date, disciplines)
                                VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                             (self.name_, datetime.now().isoformat(), self.version,
                              self.lane_style, self.comp_type, self.comp_file,
                              self.start_date_, self.end_date_, self.disciplines_))
        sponsor_img_data = ""
        if self.sponsor_img_ is not None:
            sponsor_img_data = self.sponsor_img_["data"]
        self.db_.execute('''UPDATE competition
                            SET name=?, save_date=?, version=?, lane_style=?, comp_type=?, comp_file=?,
                            start_date=?, end_date=?, sponsor_img=?, selected_country=?,
                            special_ranking_name=?, disciplines=?, aida_event_id=?,
                            aida_api_key=? WHERE id=?''',
                         (self.name_, datetime.now().isoformat(), self.version, self.lane_style, self.comp_type,
                          self.comp_file, self.start_date_, self.end_date_, sponsor_img_data,
                          self.selected_country_, self.special_ranking_name, self.disciplines_,
                          self.aida_event_id_, self.aida_api_key_,
                          self.id_))
        logging.debug("Saved competition: " + self.name)

    def load(self, comp_id):
        #TODO on load find self.id_ if not, reset to None
        load_data = self.db_.execute('''SELECT name, version, lane_style, comp_type, comp_file,
                                        start_date, end_date, sponsor_img, selected_country,
                                        special_ranking_name, publish_results,
                                        aida_event_id, aida_api_key
                                        FROM competition WHERE id=?''',
                                     comp_id)
        if load_data is None:
            logging.error("Could not find competition with id '" + str(comp_id) + "'")
            self.id_ = None
            return None
        else:
            self.id_ = comp_id
            comp_data = load_data[0]
            self.name_ = comp_data[0]
            self.version_ = comp_data[1]
            self.lane_style_ = comp_data[2]
            self.comp_type_ = comp_data[3]
            self.comp_file_ = comp_data[4]
            self.start_date_ = comp_data[5]
            self.end_date_ = comp_data[6]
            self.sponsor_img_ = {"data": comp_data[7], "aspect_ratio": 1} # TODO
            self.selected_country_ = comp_data[8]
            self.special_ranking_name_ = comp_data[9]
            self.publish_results_ = comp_data[10]
            self.aida_event_id_ = comp_data[11]
            self.aida_api_key_ = comp_data[12]
            dis = self.db_.execute('SELECT disciplines FROM block WHERE competition_id==?', (self.id_))
            self.disciplines_ = 0
            if dis is not None:
                for d in dis:
                    self.disciplines_ |= d[0]
            #self.getResultPDF('all', 'all', 'all', False, True)
            return self.name_

    def getDaysExcel(self, start_date, end_date):
        if start_date is None:
            return
        days = []
        d0 = datetime.strptime(start_date, '%Y-%m-%d')
        d1 = datetime.strptime(end_date, '%Y-%m-%d')
        delta = d1 - d0
        for i in range(delta.days + 1):
            yield (d0 + timedelta(days=i)).strftime('%Y-%m-%d')

    def getDays(self):
        days = self.db_.execute('''SELECT DISTINCT day FROM block WHERE competition_id==?''', self.id_)
        if days is None:
            return
        for day in days:
            yield u.convDay(day[0])

    def getBlocks(self):
        day_block_lane = {}
        for day in self.getDays():
            blocks = self.db_.execute('SELECT id, disciplines FROM block WHERE competition_id==? AND day==?', (self.id_, u.convDay(day)))
            if blocks is None:
                continue
            block_lane = {}
            for block in blocks:
                lanes = self.db_.execute('SELECT DISTINCT lane FROM start s WHERE s.block==? ORDER BY lane', block[0])
                if lanes is not None:
                    lanes = [self.laneStyleConverter(l[0]) for l in lanes]
                dis_s = self.disciplineIntToStr(block[1])
                block_lane[block[0]] = {'dis_s': dis_s, 'lanes': lanes}
            day_block_lane[day] = block_lane
        return day_block_lane

    def getDaysWithDisciplinesLanes(self, internal=False):
        dwd = {}
        for day in self.getDays():
            db_out = self.db_.execute('''SELECT DISTINCT discipline, lane FROM start s
                                         INNER JOIN competition_athlete ca
                                         ON s.competition_athlete_id == ca.id
                                         INNER JOIN block b
                                         ON s.block == b.id
                                         WHERE (ca.competition_id==? AND b.day==?)''',
                                      (self.id_, u.convDay(day)))
            if db_out is None:
                continue
            disciplines_on_day = list({d[0] for d in db_out})
            disciplines_w_lanes = {}
            for dis in disciplines_on_day:
                disciplines_w_lanes[dis] = [d[1] if internal else self.laneStyleConverter(d[1]) for d in db_out if d[0]==dis]
            dwd[day] = disciplines_w_lanes
        logging.debug("dwd:" + str(dwd))
        return dwd

    def getDisciplines(self):
        if self.disciplines is None:
            return []
        dwc = []
        # only aida has Overall and SpecialRanking
        if self.comp_type == "aida":
            dwc.append("Overall")
            # only add special_ranking result if we have at least one special_ranking
            has_special_ranking = self.db_.execute(
                '''SELECT id FROM competition_athlete
                WHERE competition_id==? AND special_ranking''',
                self.id_)
            if has_special_ranking is not None:
                dwc.append(self.special_ranking_name)
        dwc += self.disciplines
        return dwc

    def getCountries(self, for_result=False):
        countries = []
        if self.countries is None:
            return []
        if self.selected_country != "none":
            countries.append(self.selected_country)
        countries.append("International")
        if for_result:
            if self.comp_type == "cmas":
                # only add special_ranking result if we have at least one special_ranking
                has_special_ranking = self.db_.execute(
                    '''SELECT id FROM competition_athlete
                    WHERE competition_id==? AND special_ranking''',
                    self.id_)
                if has_special_ranking is not None:
                    countries.append(self.special_ranking_name)
        else:
            countries += self.countries
        return countries

    def getStartList(self, day, block):
        if self.comp_file is None:
            return None
        db_out = self.db_.execute('''
            SELECT a.first_name, a.last_name, a.country, s.AP, s.OT, s.lane, s.id, s.discipline, s.PB, s.dive_time
            FROM start s
            INNER JOIN competition_athlete ca ON s.competition_athlete_id == ca.id
            INNER JOIN athlete a ON ca.athlete_id == a.id
            INNER JOIN block b ON s.block == b.id
            WHERE ca.competition_id==? AND s.block==? AND b.day==?''',
            (self.id_, block, u.convDay(day)))
        if db_out is None:
            return []
        startlist = [{'Name': r[0] + " " + r[1],
                      'Nationality': r[2],
                      'AP': self.convertPerformance(r[3], r[7]),
                      'Dive Time': self.formatDiveTime(r[9]),
                      'PB': self.convertPerformance(r[8], r[7]),
                      'Warmup': self.getWTfromOT(u.convTime(r[4])),
                      'OT': u.convTime(r[4]),
                      'Lane': self.laneStyleConverter(r[5]),
                      'Discipline': r[7],
                      'Id': r[6]}
                      for r in db_out]
        startlist.sort(key=lambda r: (self.getMinFromTime(r['OT']), int(self.laneStyleConverter(r['Lane'], True))))

        br_out = self.db_.execute('''
            SELECT br.duration, br.idx
            FROM break br
            INNER JOIN block bl ON bl.id == br.block
            WHERE br.competition_id == ? AND br.block == ? AND bl.day == ?''',
            (self.id_, block, u.convDay(day)))
        if br_out is not None:
            for br in br_out:
                idx = int(br[1])
                if idx >= len(startlist):
                    continue
                br_time = str(int(br[0]/60)) + ":" + str(br[0]%60).zfill(2)
                startlist.insert(idx, {'Name': "Break", 'Nationality': "", 'AP': br_time, 'PB': '', 'Warmup': "",
                                       'OT': "", 'Lane': "", 'Discipline': "", 'Id': -1, 'Dive Time': ""})

        return startlist

    def updateStartList(self, day, block, to_remove, startlist):
        day = self.cleanDay(day)
        block = self.cleanBlock(block)
        if block is None or day == INVALID_DATE:
            return -1
        to_remove = [int(tr) for tr in to_remove]
        # apply the whole start list update atomically
        with self.db_.transaction():
            return self.updateStartListImpl(day, block, to_remove, startlist)

    def updateStartListImpl(self, day, block, to_remove, startlist):
        # remove all starts from the start list that were removed and make sure they belong to this comp
        if len(to_remove) > 0:
            placeholders = ",".join("?" * len(to_remove))
            self.db_.execute(
                '''DELETE FROM start WHERE id IN
                   (SELECT s.id FROM start s
                    INNER JOIN competition_athlete ca ON ca.id == s.competition_athlete_id
                    WHERE s.id IN ({}) AND ca.competition_id == ?)'''.format(placeholders),
                tuple(to_remove) + (self.id_, ))

        # remove all breaks
        self.db_.execute('''DELETE FROM break
                            WHERE competition_id == ? AND block == ?''',
                         (self.id_, block))

        for i in range(len(startlist)):
            if startlist[i]["Name"] == "Break":
                duration = self.getMinFromTime(self.cleanTime(startlist[i]["AP"]))
                self.db_.execute(
                    '''INSERT INTO break
                       (competition_id, block, duration, idx) VALUES (?, ?, ?, ?)''',
                    (self.id_, block, duration, i))
            else: # start
                ca_id = int(startlist[i]["Id"])
                if ca_id < 0: # new start, in this case ca_id = - athlete_id
                    ca_id = self.db_.execute(
                        '''SELECT id FROM competition_athlete
                           WHERE athlete_id == ? AND competition_id == ?''',
                        (-ca_id, self.id_))
                else: # old start in this case ca_id = start_id
                    ca_id = self.db_.execute(
                        '''SELECT s.id FROM competition_athlete ca
                           INNER JOIN start s ON ca.id == s.competition_athlete_id
                           WHERE s.id == ? AND ca.competition_id == ?''',
                        (ca_id, self.id_))
                if ca_id is None:
                    logging.warning("Invalid athlete not added to competition")
                    continue
                ot = self.cleanTime(startlist[i]["OT"])
                dive_time = self.getMinFromTime(self.cleanTime(startlist[i]["Dive Time"])) if "Dive Time" in startlist[i].keys() else 0
                discipline = self.cleanDiscipline(startlist[i]['Discipline'], block)
                if discipline is None:
                    continue
                ap = self.cleanPerf(startlist[i]["AP"], discipline)
                pb = self.cleanPerf(startlist[i]["PB"], discipline)
                if discipline == "STA":
                    ap = self.getMinFromTime(ap)
                    pb = self.getMinFromTime(pb)
                lane = self.cleanNumber(self.laneStyleConverter(startlist[i]["Lane"], True)) # TODO min/max
                if int(startlist[i]["Id"]) < 0: # new start
                    self.db_.execute(
                        '''INSERT INTO start
                           (competition_athlete_id, discipline, block, lane, OT, AP, PB, dive_time)
                           VALUES (?, ?, ?, ?, ?, ?, ?, ?)''',
                        (ca_id[0][0], discipline, block, lane, u.convTime(ot), ap, pb, dive_time))
                else: #update start
                    self.db_.execute(
                        '''UPDATE start SET
                           discipline=?, block=?, lane=?, OT=?, AP=?, PB=?, dive_time=?
                           WHERE id == ?''',
                        (discipline, block, lane, u.convTime(ot), ap, pb, dive_time, ca_id[0][0]))
        return 0

    def convertPerformance(self, val, dis):
        if val is None or val == "":
            return ""
        if dis == "STA":
            m = math.floor(int(val)/60)
            s = int(val) - m*60
            out = str(int(m)) + ":" + str(int(s)).zfill(2)
            return out
        else:
            return str(val)

    def getWTfromOT(self, ot):
        otf = self.getMinFromTime(ot)
        wtf = otf-45 # does not work if ot is close to midnight, but seriously?
        wt = str(math.floor(wtf/60)) + ":" + str(wtf%60).zfill(2)
        return wt

    def getLaneList(self, day, block, lane):
        lane_db = self.laneStyleConverter(lane, True)
        db_out = self.db_.execute('''SELECT a.first_name, a.last_name, s.AP, s.OT, a.country, a.gender, a.id, s.id, s.PB, s.discipline, s.RP, s.card, s.remarks, s.dive_time
                                     FROM athlete a
                                     INNER JOIN competition_athlete ca ON a.id == ca.athlete_id
                                     INNER JOIN start s ON s.competition_athlete_id == ca.id
                                     INNER JOIN block b ON b.id == s.block
                                     WHERE s.block == ? AND s.lane == ? AND b.day == ? AND ca.competition_id == ?
                                  ''',
                                  (block, lane_db, u.convDay(day), self.id_))
        if db_out is None:
            return -1, None
        lane_list = [{'id': r[6],
                      's_id': r[7],
                      'OT': u.convTime(r[3]),
                      'Dis': r[9],
                      'Name': r[0] + " " + r[1],
                      'Nat': r[4],
                      'AP': self.convertPerformance(r[2], r[9]),
                      'Dive Time': self.formatDiveTime(r[13]),
                      'PB': self.convertPerformance(r[8], r[9]),
                      'RP': self.convertPerformance(r[10], r[9]),
                      'Card': r[11],
                      'Remarks': r[12],
                      'NR': self.getNr(r[4], "", r[5], r[9], True)}
                       for r in db_out]
        lane_list.sort(key=lambda r: self.getMinFromTime(r['OT']))
        return 0, {'lane_list': lane_list}

    def getNr(self, country, cls, gender, discipline, convert=False):
        nr = u.NR(self.comp_type, country, cls, gender, discipline)
        if self.nr is not None and nr in self.nr:
            if convert:
                return self.convertPerformance(self.nr.get(nr), discipline)
            else:
                return self.nr.get(nr)
        else:
            return ""

    def getRecord(self, country, cls, gender, discipline, tier):
        """Record value for one tier ('NR', 'CR' or 'WR'), '' if unknown.

        CR and WR are stored per country (the AIDA API resolves continent
        membership server-side), so all tiers share the same lookup key.
        """
        if tier == "NR":
            return self.getNr(country, cls, gender, discipline)
        if self.tiered_records_ is None:
            rows = self.db_.execute(
                '''SELECT country, class, gender, discipline, tier, value
                   FROM records WHERE federation=? AND tier != 'NR' ''',
                self.comp_type)
            self.tiered_records_ = {}
            for r in (rows or []):
                self.tiered_records_[(r[0], r[1], r[2], r[3], r[4])] = r[5]
        return self.tiered_records_.get((country, cls, gender, discipline,
                                         tier), "")

    def getResult(self, discipline, gender, country, with_empty=True):
        if self.comp_file is None:
            return -1, None
        result = []
        result_keys = ['Rank', 'Name', 'Country']
        if self.comp_type == "cmas":
            result_keys.append('Club')
        if discipline == "Overall" or discipline == self.special_ranking_name:
            if self.comp_type == "cmas":
                logging.error("Attempted to get " + discipline + " ranking for cmas competition")
                return -1, None
            cmd = '''SELECT ca.id, a.first_name, a.last_name, a.country, a.club,
                            s.rp, s.penalty, s.card, s.remarks, s.discipline
                     FROM start s
                     INNER JOIN competition_athlete ca ON s.competition_athlete_id == ca.id
                     INNER JOIN athlete a ON ca.athlete_id == a.id
                     WHERE a.gender == ? AND s.remarks IS NOT NULL AND ca.competition_id == ?'''
            args = (gender, self.id_)

            if country != 'International':
                cmd += " AND a.country = ?"
                args += (country, )
                cmd += self.addEligibleCheck()
            if discipline == self.special_ranking_name:
                cmd += " AND ca.special_ranking"

            db_out = self.db_.execute(cmd, args)
            if db_out is None:
                # no results entered yet is a valid state, not an error
                # (mirrors the per-discipline branch below)
                return 0, {'results': [], 'keys': []}

            res = {}
            for r in db_out:
                if not r[0] in res:
                    res[r[0]] = {'Rank': 0, 'Name': r[1] + " " + r[2], 'Country': r[3], 'Points': 0.}
                    for d in self.disciplines:
                        res[r[0]][d] = ''
                    if self.comp_type == "cmas":
                        res[r[0]]['Club'] = r[4]
                res[r[0]][r[9]] = 0 if r[7] == "RED" else self.convertPerformance(r[5], r[9])
                res[r[0]]['Points'] += self.computePoints(r[5], r[6], r[7], r[8], r[9])
            # remove all 0 points and format points to two decimals after comma
            to_remove = []
            for r in res:
                if res[r]['Points'] == 0.:
                    to_remove.append(r)
                else:
                    res[r]['Points'] = "%.2f" % res[r]['Points']
            for tr in to_remove:
                res.pop(tr)

            res_list = sorted(list(res.values()), key=lambda r: -float(r['Points']))
            # set ranks (res_list can be empty when every entry had 0 points)
            if len(res_list) > 0:
                res_list[0]['Rank'] = 1
            for i in range(len(res_list)-1):
                if res_list[i]['Points'] != res_list[i+1]['Points']:
                    res_list[i+1]['Rank'] = i+2
                else:
                    res_list[i+1]['Rank'] = ""

            result_keys += self.disciplines + ["Points"]
            return 0, {'results': res_list, 'keys': result_keys}
        else:
            cmd = '''SELECT a.first_name, a.last_name, a.country, a.club,
                            s.AP, s.RP, s.penalty, s.card, s.remarks, s.id, s.OT, a.gender, s.judge_remarks
                     FROM start s
                     INNER JOIN competition_athlete ca ON s.competition_athlete_id == ca.id
                     INNER JOIN athlete a ON ca.athlete_id == a.id
                     WHERE s.discipline == ? AND a.gender == ? and ca.competition_id == ?'''
            if not with_empty: # remove unset results if requested
                cmd += " AND s.remarks IS NOT NULL"
            args = (discipline, gender, self.id_)
            # Add condition on country
            if country != 'International' and country != self.special_ranking_name:
                cmd += " AND a.country = ?"
                args += (country, )
                cmd += self.addEligibleCheck()
            if self.comp_type == "cmas" and country == self.special_ranking_name:
                if self.selected_country == "none":
                    cmd += " AND ca.special_ranking"
                else:
                    cmd += " AND a.country = ? AND ca.special_ranking"
                    args += (self.selected_country, )
            db_out = self.db_.execute(cmd, args)

            if db_out is None:
                return 0, {'results': [], 'keys': []}

            def check_nr(country, gender, rp, card):
                if rp is None or card != "WHITE":
                    return ""
                # flag only the highest tier the performance beats; plain
                # text, the react frontend escapes any markup
                for tier in ("WR", "CR", "NR"):
                    record = self.getRecord(country, "", gender, discipline,
                                            tier)
                    if record != "" and record < rp:
                        return ", " + tier
                return ""

            if self.comp_type == "aida":
                result = [{'Rank': i,
                           'Name': r[0] + " " + r[1],
                           'Country': r[2],
                           'AP_float': r[4],
                           'AP': self.convertPerformance(r[4], discipline),
                           'RP': self.convertPerformance(r[5], discipline) if r[8] != "DNS" else 0.,
                           'Penalty': round(r[6], 1) if r[8] != "DNS" and r[5] is not None else "",
                           'Card': r[7] if r[8] != "DNS" and r[5] is not None else "",
                           'Remarks': (("" if r[8] is None else r[8]) + check_nr(r[2], r[11], r[5], r[7])) if r[5] is not None or r[8] == "DNS" else "",
                           'JudgeRemarks': r[12],
                           'Points': ("%.2f" % self.computePoints(r[5], r[6], r[7], r[8], discipline)),
                           'Id': r[9],
                           'OT': u.convTime(r[10])}
                           for i,r in enumerate(db_out)]
                result_keys += ["AP", "RP", "Penalty", "Card", "Remarks", "Points", "JudgeRemarks"]

                # sorting according to rp (descending), card (white before others), ap (descending)
                result.sort(key=lambda r: self.sortResultsWeightsAida(r))
            else:
                result = [{'Rank': i,
                           'Name': r[0] + " " + r[1],
                           'Country': r[2],
                           'Club': r[3],
                           'RP': self.convertPerformance(r[5], discipline) if r[8] != "DNS" else 0.,
                           'Penalty': round(r[6], 1) if r[8] != "DNS" and r[5] is not None else "",
                           'Card': r[7] if r[8] != "DNS" else "",
                           'Remarks': r[8],
                           'JudgeRemarks': r[12],
                           'Points': "%.2f" % self.computePoints(r[5], r[6], r[7], r[8], discipline),
                           'Id': r[9]}
                           for i,r in enumerate(db_out)]
                result_keys += ["RP", "Card", "Remarks", "JudgeRemarks"]

                # sorting according to rp (descending)
                result.sort(key=lambda r: (-float(r['Points']), 0 if r['Remarks'] != "DNS" else 1))

            cur_points = 1000000
            cur_ap = ""
            for i in range(len(result)):
                if float(result[i]["Points"]) == 0.:
                    result[i]["Rank"] = ""
                elif result[i]["Points"] == cur_points and (self.comp_type == "cmas" or result[i]["AP"] == cur_ap):
                    result[i]["Rank"] = ""
                else:
                    cur_points = result[i]["Points"]
                    if self.comp_type == "aida":
                        cur_ap = result[i]["AP"]
                    result[i]["Rank"] = i+1
            return 0, {'results': result, 'keys': result_keys}

    def sortResultsWeightsAida(self, r):
        w0 = -float(r['Points'])
        w1 = 0
        w2 = -r['AP_float']
        if r['Card'] != "WHITE":
            w1 = 1
        if r['RP'] == "" or r['Remarks'] == "DNS":
            w1 = 2
            w2 = self.getMinFromTime(r['OT'])*100 + r['AP_float']
        return (w0, w1, w2)

    def computePoints(self, rp, penalty, card, remarks, discipline):
        if card == "RED" or remarks == "DNS" or rp is None:
            return 0.
        if discipline == "STA":
            return max(int(rp)*0.2 - penalty, 0.0)
        elif discipline in POOL_DISCIPLINES:
            if self.comp_type == "aida":
                return max(int(rp)*0.5 - penalty, 0.0)
            else:
                return max(int(rp*2.)*0.25 - penalty, 0.0)
        else:
            return max(int(rp) - penalty, 0.0)

    def changeSelectedCountry(self, country):
        if country == "none":
            self.selected_country_ = None
        else:
            if country in self.countries:
                self.selected_country_ = country
            else:
                return 1
        self.save()
        return 0

    def getBreaks(self, day):
        min_break = 24*60
        breaks_list = []
        found = False
        for d in self.getDays():
            if d == day:
                found = True
                break
        if not found:
            return None
        athletes = self.db_.execute(
            '''SELECT a.id, a.first_name, a.last_name FROM athlete a
               INNER JOIN competition_athlete ca ON ca.athlete_id == a.id
               WHERE ca.competition_id == ?''',
            self.id_)
        if athletes is None:
            return None
        for a in athletes:
            db_out = self.db_.execute(
                '''SELECT s.OT, s.discipline FROM start s
                   INNER JOIN competition_athlete ca ON s.competition_athlete_id == ca.id
                   INNER JOIN block b ON b.id == s.block
                   WHERE b.day = ? AND ca.competition_id = ? AND ca.athlete_id == ?''',
                (u.convDay(day), self.id_, a[0]))
            if db_out is None:
                continue
            n = len(db_out)
            for i in range(n-1):
                this_break = {"Name": a[1] + " " + a[2]}
                this_break["Dis1"] = db_out[i][1]
                this_break["Dis2"] = db_out[i+1][1]
                this_break["OT1"] = u.convTime(db_out[i][0])
                this_break["OT2"] = u.convTime(db_out[i+1][0])
                time = self.getMinFromTime(this_break["OT2"]) - self.getMinFromTime(this_break["OT1"])
                if time < 0:
                    time = -time
                    tmp = this_break["OT1"]
                    this_break["OT1"] = this_break["OT2"]
                    this_break["OT2"] = tmp
                    tmp = this_break["Dis1"]
                    this_break["Dis1"] = this_break["Dis2"]
                    this_break["Dis2"] = tmp
                min_break = min(min_break, time)
                this_break["Break"] = str(int(time/60)).zfill(2) + ":" + str(time%60).zfill(2)
                breaks_list.append(this_break)
        min_break = str(int(min_break/60)).zfill(2) + ":" + str(min_break%60).zfill(2)
        breaks_list = sorted(breaks_list, key=lambda d: d["Break"])
        return {"min_break": min_break, "breaks_list": breaks_list}

    def setOTs(self, data):
        ots = []
        for day in self.getDays():
            ots_on_day = self.db_.execute(
                '''SELECT DISTINCT s.OT from start s
                   INNER JOIN competition_athlete ca ON ca.id == s.competition_athlete_id
                   INNER JOIN block b ON b.id == s.block
                   WHERE ca.competition_id == ? AND b.day == ?''',
                (self.id_, u.convDay(day)))
            if ots_on_day is None:
                continue
            dayc = ":".join(day.split("-"))
            ots += [dayc + ":" + u.convTime(ot[0]) + ":00" for ot in ots_on_day]
        data["ots"] = ots

    def setSpecialRankingName(self, data):
        data["special_ranking_name"] = self.special_ranking_name

    def changeSpecialRankingName(self, name):
        self.special_ranking_name_ = name
        self.save()

    def parseTime(self, date):
        if type(date) is time:
            return date.strftime('%H:%M')
        elif type(date) is str:
            return date
        else:
            return str(date)

    def formatDiveTime(self, dive_time_seconds):
        minutes = dive_time_seconds // 60
        seconds = dive_time_seconds - minutes*60
        return str(minutes) + ":" + str(seconds).zfill(2)

    def formatSTA(self, minutes, seconds):
        if seconds == "" or math.isnan(seconds):
            return ""
        out = str(int(minutes)) + ":"
        if self.comp_type == "aida":
            out += str(int(seconds)).zfill(2)
        else:
            out += "%05.2f" % round(seconds, 2)
        return out

    # returns tuple with ca_id and if it athlete belongs to any other comp
    def isAthleteInCompetition(self, athlete_id):
        db_out = self.db_.execute(
            "SELECT id FROM competition_athlete WHERE competition_id==? AND athlete_id==?",
            (self.id_, athlete_id))
        if db_out is None:
            return None, None
        else:
            db_comps = self.db_.execute(
                "SELECT id FROM competition_athlete WHERE athlete_id==?",
                athlete_id)
            in_other_comp = len(db_comps) > 1
            return db_out[0][0], in_other_comp

    def deleteAthlete(self, ca_id, a_id, in_other_comp):
        with self.db_.transaction():
            self.db_.execute("DELETE FROM start WHERE competition_athlete_id=?", ca_id)
            self.db_.execute("DELETE FROM competition_athlete WHERE id=?", ca_id)
            if not in_other_comp:
                self.db_.execute("DELETE FROM athlete WHERE id=?", a_id)

    def getAthleteData(self, data):
        data["athletes"] = []
        athletes = self.db_.execute(
            '''SELECT a.id, a.first_name, a.last_name, a.gender, a.country, a.aida_id, a.club,
                      ca.special_ranking, ca.paid, ca.medical_checked, ca.registered, ca.eligible_national
               FROM athlete a
               INNER JOIN competition_athlete ca ON ca.athlete_id == a.id
               WHERE ca.competition_id == ?''',
            self.id_)
        if athletes is None:
            return
        for a in athletes:
            data["athletes"].append(
                {"last_name": a[2], "first_name": a[1], "gender": a[3],
                 "country": a[4], "id": a[0], "aida_id": a[5],
                 "club": a[6], "special_ranking": a[7], "paid": a[8],
                 "medical_checked": a[9], "registered": a[10],
                 "eligible_national": a[11]})

    def addAthlete(self, first_name, last_name, gender, country, club, aida_id):
        country = self.cleanCountry(country)
        if country is None:
            return -1
        a_id = self.db_.execute(
            "SELECT id FROM athlete WHERE first_name == ? AND last_name == ? AND country == ?",
            (first_name, last_name, country))
        a = None
        # user added in other comp
        if a_id is not None:
            ca_id = self.db_.execute(
                "SELECT id FROM competition_athlete WHERE athlete_id == ? AND competition_id == ?",
                (a_id[0][0], self.id_))
            if ca_id is not None:
                return 1 # athlete already exists in current competition
            else:
                # not in current competition, so add it
                self.db_.execute(
                    '''INSERT INTO competition_athlete
                       (competition_id, athlete_id, special_ranking) VALUES(?, ?, ?)''',
                    (self.id_, a_id[0][0], False))
                # TODO check if rest of arguments are identical
                return 0
        else:
            a = athlete.Athlete.fromArgs(aida_id, first_name, last_name, gender, country, club, self.db_)
            a.associateWithComp(self.id_)
            return 0

    def cleanDiscipline(self, dis, block=None):
        # for a list, we check each element
        if isinstance(dis, list):
            for d in dis:
                res = self.cleanDiscipline(d, block)
                if res is None:
                    return None
            return dis

        if not self.isMainDiscipline(dis):
            return None
        else:
            if block is not None:
                dis_i = DISCIPLINES.index(dis)
                dis_in_block = self.db_.execute('SELECT disciplines FROM block WHERE id==?', block)
                if dis_in_block is None or (dis_in_block[0][0] & 1<<dis_i) == 0:
                    return None
            return dis

    def cleanCountry(self, country):
        country = regex.compile('[A-Z][A-Z][A-Z]$').match(str(country).strip())
        if country is None:
            return None
        else:
            return country[0]

    def cleanTime(self, time):
        ctime = regex.compile('[0-5]?\\d:[0-5]\\d$').match(str(time))
        if ctime is None:
            return INVALID_TIME
        else:
            return ctime[0]

    def cleanNumber(self, n, digits = 0, minval = -sys.float_info.max, maxval = sys.float_info.max):
        pstr = '^\\d+'
        if digits > 0:
            pstr += '.?'
        pstr += ''.join(['\\d?' for i in range(digits)])
        pattern = regex.compile(pstr)
        cn = pattern.match(str(n))
        if cn is None:
            return "0" if minval == -sys.float_info.max else str(minval)
        elif minval <= float(cn[0]) and float(cn[0]) <= maxval:
            return cn[0]
        else:
            return str(minval)

    def cleanPerf(self, perf, discipline):
        if discipline == "STA":
            return self.cleanTime(perf)
        elif self.comp_type == "aida" or discipline in DEPTH_DISCIPLINES:
            return self.cleanNumber(perf, 0, 0)
        else: # cmas dynamic disciplines
            return self.cleanNumber(perf, 1, 0)

    def cleanBlock(self, block):
        block = int(block)
        block = self.db_.execute('SELECT id FROM block WHERE id==? AND competition_id==?', (block, self.id_))
        if block is not None:
            return block[0][0]
        else:
            return None

    def cleanDay(self, day, format_only = False):
        if format_only:
            day = str(day)
            res = True
            try:
                res = bool(datetime.strptime(day, "%Y-%m-%d"))
            except ValueError:
                res = False
            return day if res else INVALID_DATE
        else:
            if day in [d for d in self.getDays()]:
                return str(day)
            else:
                return INVALID_DATE

    def cleanCard(self, card):
        if card in ["WHITE", "YELLOW", "RED"]:
            return card
        return None

    def cleanPenalty(self, penalty):
        try:
            return float(penalty)
        except:
            return 0.

    def getMinFromTime(self, time_str):
        h_m = time_str.split(':')
        return int(h_m[0])*60 + int(h_m[1])

    def updateResult(self, s_id, rp, penalty, card, remarks, judge_remarks):
        s_id = int(s_id)
        penalty = self.cleanPenalty(penalty)
        card = self.cleanCard(card)
        if card is None:
            return 1, None
        # TODO sanity checks for card and remarks and discipline
        ap_dis = self.db_.execute(
                '''SELECT s.ap, s.discipline
                   FROM competition_athlete ca
                   INNER JOIN start s ON ca.id == s.competition_athlete_id
                   WHERE s.id == ?''',
                s_id)
        if ap_dis is not None:
            discipline = ap_dis[0][1]
            rp = self.cleanPerf(rp, discipline)
            if discipline == "STA":
                rp = self.getMinFromTime(rp)
            under_ap_penalty = self.getUnderApPenalty(ap_dis[0][0], rp, discipline, card) if self.comp_type == "aida" else 0
            self.db_.execute(
                '''UPDATE start SET rp = ?, penalty = ?, card = ?, remarks = ?, judge_remarks = ? WHERE id == ?''',
                (rp, under_ap_penalty + penalty, card, remarks, judge_remarks, s_id))
            return 0, None
        return 1, None

    def getUnderApPenalty(self, ap, rp, discipline, card):
        if card != "YELLOW":
            return 0
        factor = 1.
        if discipline == "STA":
            factor = 0.2
        elif discipline[0] == "D":
            factor = 0.5
        rp = float(rp)
        if rp < ap:
            return round((ap - rp)*factor, 1)
        else:
            return 0.

    def updateNationalRecords(self, client=None):
        """Refresh the AIDA records used for the record flags.

        With an API key configured the records come from the AIDA API (all
        tiers NR/CR/WR, scoped to this competition's countries and
        disciplines); without one the public records page is scraped as
        before (NR only, all countries).
        """
        # Do check if country converter is >= 1.2
        try:
            if self.has_aida_api_key:
                records, countries = self.fetchAidaRecords(client)
            else:
                records = [(key.country, key.cls, key.gender, key.discipline,
                            "NR", value)
                           for key, value in u.getNationalRecordsAida().items()]
                countries = None  # the scraper covers all countries
            with self.db_.transaction():
                if countries is None:
                    self.db_.execute(
                        "DELETE FROM records WHERE federation='aida'")
                else:
                    # the records table is shared between competitions, so
                    # only replace the rows of the countries just queried
                    for country in countries:
                        self.db_.execute('''DELETE FROM records
                                         WHERE federation='aida' AND country=?''',
                                         country)
                for country, cls, gender, discipline, tier, value in records:
                    self.db_.execute('''INSERT INTO records
                                     ('federation', 'country', 'class', 'gender',
                                      'discipline', 'value', 'tier')
                                     VALUES ('aida', ?, ?, ?, ?, ?, ?)''',
                                     (country, cls, gender, discipline, value,
                                      tier))
            self.nrs_ = None  # reload lazily from the updated table
            self.tiered_records_ = None
        except AidaApiError as e:
            # the client's messages are safe to show to the admin
            return 1, {"error_msg": str(e)}
        except Exception as e:
            logging.warning("Updating national records failed: %s", e)
            return 1, None
        return 0, None

    def fetchAidaRecords(self, client=None):
        """Fetch NR/CR/WR from the AIDA API for this competition's field.

        Queries one (country, discipline, gender) combination at a time,
        limited to the countries of the registered athletes and the
        competition's disciplines. Returns (rows for the records table,
        IOC codes of the countries that were queried).
        """
        client = client if client is not None else self.aidaClient()
        rows = self.db_.execute(
            '''SELECT DISTINCT a.country FROM athlete a
               INNER JOIN competition_athlete ca ON ca.athlete_id == a.id
               WHERE ca.competition_id == ?''', self.id_)
        cc = self.countryConverter()
        records = []
        queried = []
        for row in (rows or []):
            ioc = row[0]
            # the records endpoint wants 2-letter codes, Compy stores IOC
            abrv = cc.convert(names=str(ioc), src="IOC", to="ISO2",
                              not_found=None)
            if not isinstance(abrv, str) or len(abrv) != 2:
                logging.warning("No ISO2 code for country %s, "
                                "skipping its records", ioc)
                continue
            queried.append(ioc)
            for discipline in self.disciplines:
                for gender in ("M", "F"):
                    payload = client.getRecords(abrv, discipline, gender)
                    for key, tier in (("nr", "NR"), ("cr", "CR"),
                                      ("wr", "WR")):
                        # assumption (unverified against the live API): the
                        # first list entry is the current record
                        entries = payload[key]
                        if len(entries) == 0 \
                           or not isinstance(entries[0], dict):
                            continue
                        value = self.aidaPerformance(
                            entries[0].get("reportedPerformance"),
                            entries[0].get("performanceMeasurement"),
                            discipline)
                        if value is not None:
                            records.append((ioc, "", gender, discipline,
                                            tier, value))
        return records, queried

    def isJudgeInCompetition(self, judge_id):
        db_out = self.db_.execute(
            "SELECT id FROM judge WHERE competition_id==? AND id==?",
            (self.id_, judge_id))
        if db_out is None:
            return None
        else:
            return db_out[0][0]

    def deleteJudge(self, judge_id):
        self.db_.execute("DELETE FROM judge WHERE id=?", judge_id)

    def getJudgeData(self, data):
        data["judges"] = []
        judges = self.db_.execute(
            '''SELECT id, first_name, last_name FROM judge WHERE competition_id == ?''',
            self.id_)
        if judges is None:
            return
        for j in judges:
            data["judges"].append(
                {"last_name": j[2], "first_name": j[1], "id": j[0]})

    def addJudge(self, first_name, last_name):
        a_id = self.db_.execute(
            "SELECT id FROM judge WHERE first_name == ? AND last_name == ? AND competition_id == ?",
            (first_name, last_name, self.id_))
        a = None
        # user already in this comp
        if a_id is not None:
            return 1;
        else:
            salt = random.randrange(sys.maxsize)
            self.db_.insert(
                "INSERT INTO judge (first_name, last_name, salt, competition_id) VALUES (?, ?, ?, ?)",
                (first_name, last_name, salt, self.id_))
            return 0

    def getJudgeQrCode(self, judge_id, url_root):
        db_out = self.db_.execute(
            "SELECT first_name, last_name, salt FROM judge WHERE competition_id==? AND id==?",
            (self.id_, judge_id))
        if db_out is None:
            return None
        else:
            jhash = self.getJudgeQrCodeHash(db_out[0][0], db_out[0][1], db_out[0][2], judge_id)
            # TODO make configurable
            url = url_root + "judge/" + str(self.id_) + "/" + str(judge_id)+"?hash=" + jhash
            qr_code = qrcode.make(url)
            ba = BytesIO()
            qr_code.save(ba, format='PNG')
            qr_code_base64 = "data:image/png;base64," + base64.b64encode(ba.getvalue()).decode('utf-8')
            return (qr_code_base64, db_out[0][0], db_out[0][1], url)

    def getJudgeQrCodeHash(self, first_name, last_name, salt, judge_id):
        return hashlib.sha256((first_name + last_name + str(salt) + str(judge_id)).encode('utf-8')).hexdigest()

    def getCompDataAndValidateJudge(self, comp_id, judge_id, judge_hash):
        try:
            assert(comp_id is not None)
            assert(judge_id is not None)
            comp_id = int(comp_id)
            judge_id = int(judge_id)
            assert(judge_hash.isalnum())
        except:
            return -1, None
        db_out = self.db_.execute(
            "SELECT first_name, last_name, salt FROM judge WHERE competition_id==? AND id==?",
            (comp_id, judge_id))
        db_out2 = self.db_.execute("SELECT name, comp_type FROM competition WHERE id==?", comp_id)
        if db_out is None or db_out2 is None:
            return -1, None

        first_name = db_out[0][0]
        last_name = db_out[0][1]
        judge_hash_db = self.getJudgeQrCodeHash(first_name, last_name, db_out[0][2], judge_id)
        if judge_hash_db != judge_hash:
            # do not touch any state for a failed validation
            return -1, None

        self.id_ = comp_id
        return 0, {'comp_name': db_out2[0][0],
                   'first_name': first_name,
                   'last_name': last_name,
                   'federation': db_out2[0][1]}

    def getAthleteResult(self, s_id):
        try:
            s_id = int(s_id)
        except:
            return -1, None

        db_out = self.db_.execute(
            '''SELECT a.first_name, a.last_name, a.country, s.AP, s.RP, s.penalty, s.card,
                      s.remarks, s.OT, a.gender, s.judge_remarks, s.discipline, s.PB,
                      b.day, s.block, s.lane
               FROM start s
               INNER JOIN competition_athlete ca ON s.competition_athlete_id == ca.id
               INNER JOIN athlete a ON ca.athlete_id == a.id
               INNER JOIN block b ON b.id == s.block
               WHERE ca.competition_id == ? AND s.id == ?''',
            (self.id_, s_id))

        if db_out is None:
            return -1, None

        discipline = db_out[0][11]
        ret, lane_list_dict = self.getLaneList(u.convDay(db_out[0][13]), db_out[0][14], self.laneStyleConverter(db_out[0][15]))
        return 0, {'Name': db_out[0][0] + " " + db_out[0][1],
                   'Country': db_out[0][2],
                   'AP': self.convertPerformance(db_out[0][3], discipline),
                   'Dis': discipline,
                   'PB': self.convertPerformance(db_out[0][12], discipline),
                   'RP': self.convertPerformance(db_out[0][4], discipline),
                   'Penalty': db_out[0][5],
                   'Card': db_out[0][6],
                   'Remarks': db_out[0][7],
                   'JudgeRemarks': db_out[0][10],
                   'Id': s_id,
                   'OT': u.convTime(db_out[0][8]),
                   'NR': self.getNr(db_out[0][2], "", db_out[0][9], discipline, True),
                   'Gender': db_out[0][9],
                   'lane_list': lane_list_dict['lane_list']}

    def cleanFederation(self, federation):
        if not federation in FEDERATIONS:
            return None
        else:
            return federation

    def getAllDisciplines(self, federation):
        federation = self.cleanFederation(federation)
        if federation is None:
            return None
        else:
            return DISCIPLINES

    def modifyBlock(self, day, disciplines, block, add):
        day = self.cleanDay(day, True)
        disciplines = self.cleanDiscipline(disciplines)
        dis_i = self.disciplineListToInt(disciplines)
        if add:
            # check if entry exists already, if yes, we don't allow another
            db_out = self.db_.execute('SELECT id FROM block WHERE competition_id==? AND day==? AND disciplines==?',
                                      (self.id_, u.convDay(day), dis_i))
            if db_out is None:
                self.db_.execute('INSERT INTO block (competition_id, day, disciplines) VALUES (?, ?, ?)',
                                 (self.id_, u.convDay(day), dis_i))
                return 0
            else:
                return 1
        else:
            block = self.cleanBlock(block)
            if block is None:
                return 2
            self.db_.execute(
                '''UPDATE block SET day=?, disciplines=?
                   WHERE competition_id==? AND id==?''',
                (u.convDay(day), dis_i, self.id_, block))
            return 0

    def removeBlock(self, block):
        block = self.cleanBlock(block)
        if block is None:
            return 1
        self.db_.execute('DELETE FROM block WHERE id==? AND competition_id==?', (block, self.id_))
        return 0

    def getDisciplinesFromInt(self, dis_i):
        dis = []
        for bit, d in enumerate(DISCIPLINES):
            if dis_i & 1<<bit:
                dis.append(d)
        return dis

    def disciplineIntToStr(self, dis_i):
        return ", ".join(self.getDisciplinesFromInt(dis_i))

    def disciplineListToInt(self, dis):
        return sum([1<<DISCIPLINES.index(d) for d in dis])

    def getFourStarts(self, current, offset):
        comp = "<=" if current else ">"
        order = "DESC" if current else "ASC"
        cmd = '''SELECT a.first_name, a.last_name, a.country, s.OT, s.lane, s.remarks
                 FROM start s
                 INNER JOIN competition_athlete ca
                 ON s.competition_athlete_id == ca.id
                 INNER JOIN athlete a
                 ON a.id = ca.athlete_id
                 INNER JOIN block b
                 ON b.id == s.block
                 WHERE b.day*10000+s.OT == (
                    SELECT b.day*10000+s.OT
                    FROM start s
                    INNER JOIN competition_athlete ca
                    ON s.competition_athlete_id == ca.id
                    INNER JOIN athlete a
                    ON a.id = ca.athlete_id
                    INNER JOIN block b
                    ON b.id == s.block
                    WHERE ca.competition_id = ? AND b.day*10000+s.OT {} ?
                    ORDER BY b.day {}, CAST(s.OT as INTEGER) {}
                    LIMIT 1)
                 ORDER BY s.lane
                 LIMIT 4'''.format(comp, order, order)
        min_shift = 3 if self.comp_type == "cmas" else 2
        now = datetime.now(timezone.utc) + timedelta(minutes=min_shift) + timedelta(milliseconds=offset)
        today = now.year*10000 + now.month*100 + now.day
        time = now.hour*100 + now.minute
        db_out = self.db_.execute(cmd, (self.id_, today*10000 + time))
        if db_out is None:
            return None
        else:
            return [{'name': d[0] + " " + d[1],
                     'country': d[2],
                     'OT': u.convTime(d[3]),
                     'lane': self.laneStyleConverter(d[4]),
                     'dns': d[5] == "DNS"}
                     for d in db_out]

    def updatePublishResults(self, publish_results):
        if self.id_ is None:
            return 1, None
        # clean received data
        publish_results = 1 if publish_results == True else 0
        self.db_.execute("UPDATE competition SET publish_results=?  WHERE id==?", (publish_results, self.id_))
        self.publish_results_ = publish_results == 1
        return 0, None

    def getResultContent(self):
        content = {"version": self.version}
        # if no comp id is set get a list of all competitions
        if self.id_ is None:
            db_out = self.db_.execute("SELECT id, name FROM competition WHERE publish_results==1")
            if db_out is not None:
                content['data'] = {'comp_list': [{'id': d[0], 'name': d[1]} for d in db_out]}
        # if a comp id is set then get all disciplines and countries
        else:
            data = {}
            data['disciplines'] = self.getDisciplines()
            data['countries'] = self.getCountries(True)
            data['comp_name'] = self.name
            data['comp_id'] = self.id_
            content['data'] = data

        return 0, content

    def cleanCompIdPublished(self, comp_id):
        try:
            comp_id = int(comp_id)
            db_out = self.db_.execute("SELECT name FROM competition WHERE publish_results==1 AND id==?", comp_id)
            if db_out is None:
                return None, None
            return comp_id, db_out[0][0]
        except:
            return None, None

    def getResultList(self, discipline_id, gender, country_id):
        if self.id_ is None:
            return -1, None
        try:
            discipline_id = int(discipline_id)
            country_id = int(country_id)
            discipline = self.getDisciplines()[discipline_id]
            country = self.getCountries(True)[country_id]
            if not (gender == "Female" or gender == "Male"):
                return -1, None
            gender = "F" if gender == "Female" else "M"
            ret, content = self.getResult(discipline, gender, country, False)
            keys = content['keys']
            results = content['results']
            value_key = 'RP' if 'RP' in keys else 'Points'
            if not self.isMainDiscipline(discipline):
                def getIndividualResults(r):
                    indiv_results = []
                    for d in DISCIPLINES:
                        if d in keys and r[d] != "":
                            indiv_results.append({'dis': d, 'rp': r[d]})
                    return indiv_results

                return 0, {'results': [{'rank': r['Rank'],
                                        'name': r['Name'],
                                        'value': "DNS" if 'Remarks' in keys and r['Remarks'] == "DNS" else r[value_key],
                                        'individual_results': getIndividualResults(r)
                                       } for r in results]}
            else:
                return 0, {'results': [{'rank': r['Rank'],
                                        'name': r['Name'],
                                        'value': "DNS" if 'Remarks' in keys and r['Remarks'] == "DNS" else r[value_key],
                                        'country': r['Country'],
                                        'card': r['Card'] if 'Card' in keys else None,
                                        'ap': r['AP'] if 'AP' in keys else None,
                                        'penalty': r['Penalty'],
                                        'remarks': r['Remarks'] if 'Remarks' in keys else None,
                                        'points': r['Points'] if 'Points' in keys else None
                                       } for r in results]}
        except Exception as e:
            logging.debug("Failed to get result list: %s", e)
            return -1, None

    def isMainDiscipline(self, discipline):
        return discipline in DISCIPLINES

    def deleteComp(self, comp_id):
        try:
            assert(comp_id is not None)
            comp_id = int(comp_id)
        except:
            return 1, {'status': 'error', 'status_msg': 'Invalid id'}
        valid_id = self.db_.execute("SELECT COUNT(*) FROM competition WHERE id=?", comp_id)
        if valid_id is None:
            return 1, {'status': 'error', 'status_msg': 'Competition id ' + str(comp_id) + ' does not exist'}
        if comp_id == 1:
            return 1, {'status': 'error', 'status_msg': 'Cannot delete competition with id 1'}
        data = {}

        # delete the competition and all its dependent rows atomically
        with self.db_.transaction():
            self.db_.execute("DELETE FROM competition WHERE id=?", comp_id)
            self.db_.execute("DELETE FROM judge WHERE competition_id=?", comp_id)
            self.db_.execute("DELETE FROM block WHERE competition_id=?", comp_id)
            self.db_.execute("DELETE FROM break WHERE competition_id=?", comp_id)

            # Get all athletes from comp
            db_out = self.db_.execute("SELECT id, athlete_id FROM competition_athlete WHERE competition_id=?", comp_id)

            if db_out is not None:
                for ca in db_out:
                    ca_id = ca[0]
                    self.db_.execute("DELETE FROM start WHERE competition_athlete_id=?", ca_id)

                    athlete_id = ca[1]
                    n_comp_athlete = self.db_.execute("SELECT COUNT(*) FROM competition_athlete WHERE athlete_id=?", athlete_id)

                    # Delete if athlete is only in one comp (i.e. the one being deleted)
                    if n_comp_athlete is not None and n_comp_athlete[0][0] == 1:
                        self.db_.execute("DELETE FROM athlete WHERE id=?", athlete_id)

                self.db_.execute("DELETE FROM competition_athlete WHERE competition_id=?", comp_id)

        data['competitions'] = self.getSavedCompetitions()
        data['status'] = 'success'
        data['status_msg'] = 'Successfully deleted comp with id ' + str(comp_id)
        return 0, data

    def storeResults(self, fpath):
        outfile = os.path.join(self.config.download_folder, "Results_" + os.path.basename(fpath))
        with pd.ExcelWriter(outfile) as xls_writer:
            # copy first three sheets
            copy_sheets = ['Event', 'Athletes and Judges', 'Settings']
            for c in copy_sheets:
                df = pd.read_excel(fpath, sheet_name=c)
                df.to_excel(xls_writer, sheet_name=c, index=False)

            for day in self.getDays():
                df = pd.read_excel(fpath, sheet_name=day, skiprows=1)
                selector = list(zip(df['Diver Id'], df['Discipline']))
                selector_str = ", ".join(["(\"" + i + "\", \"" + d + "\")" for (i, d) in selector])
                cmd = f'''SELECT s.RP, s.penalty, s.card, s.remarks, s.AP, a.aida_id, s.discipline
                          FROM start s
                          INNER JOIN competition_athlete ca ON s.competition_athlete_id == ca.id
                          INNER JOIN athlete a ON ca.athlete_id == a.id
                          WHERE (a.aida_id, s.discipline) in (VALUES {selector_str}) AND ca.competition_id == ?'''
                db_out = self.db_.execute(cmd, self.id_)
                if db_out is None:
                    return -1, None
                df['Card'] = df['Card'].astype(str)
                df['Remarks'] = df['Remarks'].astype(str)
                for i, s in enumerate(selector):
                    j = -1
                    for jj, db_row in enumerate(db_out):
                        if db_row[5] == s[0] and db_row[6] == s[1]:
                            j = jj
                            break
                    if j == -1:
                        df.at[i, 'Card'] = ''
                        df.at[i, 'Remarks'] = 'DNS'
                        continue
                    db_row = db_out[j]
                    if db_row[0] is None:
                        # start exists but has no result recorded yet
                        df.at[i, 'Card'] = ''
                        df.at[i, 'Remarks'] = db_row[3] if db_row[3] is not None else ''
                        db_out.pop(j)
                        continue
                    if s[1] == "STA":
                        rp_min = int(db_row[0] / 60)
                        df.at[i, 'Meters or Min.1'] = rp_min
                        df.at[i, 'Sec(STA only).1'] = db_row[0] - 60*rp_min
                    else:
                        df.at[i, 'Meters or Min.1'] = db_row[0]
                    under_ap_penalty = self.getUnderApPenalty(db_row[4], db_row[0], s[1], db_row[2])
                    df.at[i, 'Pen(other)'] = max(db_row[1] - under_ap_penalty, 0)
                    if db_row[3] == 'DNS':
                        df.at[i, 'Card'] = ''
                    else:
                        df.at[i, 'Card'] = db_row[2]
                    df.at[i, 'Remarks'] = db_row[3]
                    db_out.pop(j)
                df.fillna('', inplace=True)
                df.to_excel(xls_writer, sheet_name=day, index=False)
        wb = load_workbook(outfile)
        for day in self.getDays():
            ws = wb[day]
            ws.insert_rows(1)
            ws['F1'] = 'AP'
            ws['H1'] = 'RP'
            ws['J1'] = 'Penalties'
        wb.save(outfile)
        return 0, outfile

    def addEligibleCheck(self):
        # check if at least one eligible
        has_eligible = self.db_.execute("SELECT COUNT(*) FROM competition_athlete WHERE competition_id==? AND eligible_national==1", self.id_)
        if has_eligible is not None and has_eligible[0][0] != 0:
            return " AND ca.eligible_national == 1"
        return ""
