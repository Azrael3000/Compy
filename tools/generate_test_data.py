#!/usr/bin/python3
#
#  Generates an AIDA-International-style competition Excel file that Compy
#  can import via "Refresh data" (or the /upload_file endpoint).
#
#  Usage: python3 tools/generate_test_data.py [output.xlsx]
#
#  The layout mirrors a real "Event Result Sheet" export from
#  aidainternational.org:
#    - sheet "Event":               event metadata as label/value rows
#    - sheet "Athletes and Judges": athletes (A-E) and judges (F-G) side by side
#    - sheet "Settings":            available remarks and penalty cards
#    - one sheet per competition day, columns A-T:
#        Discipline | Diver Name | Diver Id | Gender | Diver Country |
#        AP (Meters or Min / Sec(STA only)) | RP (dito) |
#        Pen(UNDER AP) | Pen(other) | Card | Remarks | Points |
#        OT | WT | Zone | Dive Time | Descend Time | Record
#
#  Differences to a real export, on purpose:
#    - "Diver Name", "Pen(UNDER AP)" and "Points" contain plain values.
#      Real exports have formulas there (openpyxl cannot store the cached
#      results pandas needs, and Compy reads values, not formulas).
#      "Diver Name" carries the athlete-profile hyperlink like the original.
#    - Days that are already over have results filled in; a real pre-event
#      export has none. The last day is "today" and left open.
#
#  Produces a 3-day competition with 30 athletes (STA+DYN, DYNB+DNF, CWT),
#  results and cards for past days, DNS entries, and an open final day.

import random
import sys
import uuid
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font

random.seed(42)

PROFILE_URL = "https://www.aidainternational.org/Athletes/Profile-"

ATHLETES = [
    # first, last, gender, country
    ("Anna",     "Berger",      "F", "AUT"),
    ("Lukas",    "Steiner",     "M", "AUT"),
    ("Marie",    "Novak",       "F", "CZE"),
    ("Tomas",    "Dvorak",      "M", "CZE"),
    ("Jana",     "Kovacova",    "F", "SVK"),
    ("Peter",    "Molnar",      "M", "SVK"),
    ("Clara",    "Schmidt",     "F", "GER"),
    ("Felix",    "Wagner",      "M", "GER"),
    ("Sophie",   "Mueller",     "F", "GER"),
    ("Jonas",    "Fischer",     "M", "GER"),
    ("Chiara",   "Rossi",       "F", "ITA"),
    ("Marco",    "Bianchi",     "M", "ITA"),
    ("Elena",    "Ricci",       "F", "ITA"),
    ("Luca",     "Moretti",     "M", "ITA"),
    ("Camille",  "Laurent",     "F", "FRA"),
    ("Hugo",     "Moreau",      "M", "FRA"),
    ("Lea",      "Dubois",      "F", "FRA"),
    ("Nathan",   "Petit",       "M", "FRA"),
    ("Ivana",    "Horvat",      "F", "CRO"),
    ("Ante",     "Kovacevic",   "M", "CRO"),
    ("Maja",     "Zupan",       "F", "SLO"),
    ("Luka",     "Kranjc",      "M", "SLO"),
    ("Zofia",    "Kowalska",    "F", "POL"),
    ("Jakub",    "Nowak",       "M", "POL"),
    ("Carmen",   "Garcia",      "F", "ESP"),
    ("Diego",    "Martinez",    "M", "ESP"),
    ("Emma",     "Jansen",      "F", "NED"),
    ("Daan",     "Visser",      "M", "NED"),
    ("Nora",     "Haugen",      "F", "NOR"),
    ("Erik",     "Berg",        "M", "NOR"),
]

JUDGES = [
    # last, first
    ("Kane",      "Arthur"),
    ("Darthmore", "Nicole"),
    ("Smith",     "Bernhard"),
    ("Burns",     "Roger"),
]

# disciplines per day; lanes available
DAY_PLAN = [
    (["STA", "DYN"],  4),   # day 1: pool
    (["DYNB", "DNF"], 4),   # day 2: pool
    (["CWT"],         4),   # day 3 (today): depth, still open
]

# remark codes as they appear on the "Settings" sheet of real exports
AVAILABLE_REMARKS = [
    "OK", "OTHER", "SHORT", "LATESTART", "GRAB", "LANYARD", "PULL", "TURN",
    "EARLYSTART", "START", "NO TAG", "UNDER AP", "DQSP", "DQJUMP", "DQOTHER",
    "DQAIRWAYS", "DQTOUCH", "DQLATESTART", "DQCHECK-IN", "DQBO-UW",
    "DQBO-SURFACE", "DQPULL", "DQOTHER-LANE", "DQPULL (MORE THAN 1X)", "DNS",
]

CARDS_WEIGHTED = ["WHITE"]*7 + ["YELLOW"]*2 + ["RED"]
REMARKS_WHITE = ["OK", "OK", "OK", "OK"]
REMARKS_YELLOW = ["TURN", "UNDER AP", "SHORT", "GRAB"]
REMARKS_RED = ["LANYARD", "DQSP", "DQBO-SURFACE", "EARLYSTART"]

DEPTH = ("CWT", "CWTB", "CNF", "FIM")


def perf_ap(dis, gender):
    """Announced performance per discipline."""
    if dis in DEPTH:
        base = random.randint(35, 85) + (5 if gender == "M" else 0)
        return base, None                       # meters
    if dis == "STA":
        m = random.randint(3, 6)
        s = random.choice([0, 15, 30, 45])
        return m, s                              # minutes, seconds
    # DYN / DYNB / DNF
    base = random.randint(50, 175) + (10 if gender == "M" else 0)
    return base, None                            # meters


def perf_rp(dis, ap_main, ap_sec, card):
    """Realized performance derived from AP and card."""
    if card == "WHITE":
        if dis == "STA":
            total = ap_main*60 + (ap_sec or 0) + random.randint(0, 40)
            return total // 60, total % 60
        return ap_main + random.randint(0, 6), None
    else:  # YELLOW (turned early / under AP), RED still has an RP
        if dis == "STA":
            total = max(60, ap_main*60 + (ap_sec or 0) - random.randint(10, 60))
            return total // 60, total % 60
        return max(10, ap_main - random.randint(1, 15)), None


def points(dis, rp_main, rp_sec, pen_under, pen_other):
    """AIDA points: 0.2/s for STA, 0.5/m for pool dynamic, 1/m for depth."""
    if dis == "STA":
        raw = (rp_main*60 + (rp_sec or 0))*0.2
    elif dis in DEPTH:
        raw = rp_main*1.
    else:
        raw = rp_main*0.5
    return round(max(raw - pen_under - pen_other, 0), 1)


def main(outfile):
    today = date.today()
    start = today - timedelta(days=2)   # two days done, today still running
    days = [start + timedelta(days=i) for i in range(len(DAY_PLAN))]

    athletes = [(str(uuid.uuid4()),) + a for a in ATHLETES]

    wb = Workbook()

    # --- Event sheet ---
    ws = wb.active
    ws.title = "Event"
    comp_name = "Compy Test Open " + str(today.year)
    for label, value in [
            ("Event Information", None),
            ("Name:", comp_name),
            ("Starts:", days[0].isoformat()),
            ("Ends:", days[-1].isoformat()),
            ("Country:", "Austria"),
            ("City:", "Testville"),
            ("Location:", "Test Pool & Lake"),
            ("Address:", "Some Road 123"),
            ("Disciplines:", ",".join(d for ds, _ in DAY_PLAN for d in ds)),
            ("Organizer:", "Compy Test Club"),
            ("Organizer Email:", "organizer@example.com"),
            ("Organizer Phone Number:", "0012345678"),
            ("Organizer Website:", None),
            ("AIDA National:", "AIDA Austria"),
            ("Competition Type:", "Pool Competition"),
            ("Additional Info:", None),
            ("Date Downloaded(UTC):", days[0].isoformat() + " 08:00:00"),
    ]:
        ws.append([label, value])

    # --- Athletes and Judges sheet (athletes A-E, judges F-G) ---
    ws = wb.create_sheet("Athletes and Judges")
    ws.append(["Athletes in Competition", None, None, None, None, "Judges"])
    ws.append(["Id", "LastName", "FirstName", "Gender", "Country",
               "LastName", "FirstName"])
    for i, (aid, first, last, gender, country) in enumerate(athletes):
        row = [aid, last, first, gender, country]
        if i < len(JUDGES):
            row += list(JUDGES[i])
        ws.append(row)

    # --- Settings sheet (remarks & cards, copied verbatim by Compy) ---
    ws = wb.create_sheet("Settings")
    ws.append(["Available Remarks", "Penalty Cards"])
    cards = ["WHITE", "YELLOW", "RED"]
    for i, remark in enumerate(AVAILABLE_REMARKS):
        ws.append([remark, cards[i] if i < len(cards) else None])

    # --- one sheet per competition day ---
    header = ["Discipline", "Diver Name", "Diver Id", "Gender",
              "Diver Country",
              "Meters or Min", "Sec(STA only)",       # AP  (F, G)
              "Meters or Min", "Sec(STA only)",       # RP  (H, I)
              "Pen(UNDER AP)", "Pen(other)",          # J, K
              "Card", "Remarks", "Points",            # L, M, N
              "OT", "WT", "Zone",                     # O, P, Q
              "Dive Time", "Descend Time", "Record"]  # R, S, T

    for day, (disciplines, n_lanes) in zip(days, DAY_PLAN):
        ws = wb.create_sheet(day.isoformat())
        ws.append([None]*5 + ["AP", None, "RP", None, "Penalties"])
        ws.append(header)
        day_done = day < date.today()

        for dis in disciplines:
            # not everybody starts in every discipline
            starters = [a for a in athletes if random.random() < 0.85]
            random.shuffle(starters)

            ot_hour, ot_min = (10, 0) if dis == disciplines[0] else (14, 0)
            lane = 1
            for aid, first, last, gender, country in starters:
                ap_main, ap_sec = perf_ap(dis, gender)
                ot = "%02d:%02d" % (ot_hour, ot_min)
                wt = "%02d:%02d" % ((ot_hour - 1, ot_min + 15) if ot_min < 45
                                    else (ot_hour, ot_min - 45))

                rp_main = rp_sec = pen_under = pen_other = card = remarks = pts = None
                dive_time = None
                if day_done:
                    if random.random() < 0.07:   # a few no-shows
                        remarks = "DNS"
                        pen_under = 0
                    else:
                        card = random.choice(CARDS_WEIGHTED)
                        rp_main, rp_sec = perf_rp(dis, ap_main, ap_sec, card)
                        pen_other = 1.0 if (card == "YELLOW" and random.random() < 0.3) else 0
                        pen_under = 0.
                        if card == "YELLOW":
                            # AIDA under-AP penalty: 0.2/s STA, 0.5/m dynamic, 1/m depth
                            if dis == "STA":
                                ap_v = ap_main*60 + (ap_sec or 0)
                                rp_v = rp_main*60 + (rp_sec or 0)
                                factor = 0.2
                            else:
                                ap_v, rp_v = ap_main, rp_main
                                factor = 0.5 if dis[0] == "D" else 1.
                            pen_under = round(max(ap_v - rp_v, 0)*factor, 1)
                        if card == "WHITE":
                            remarks = random.choice(REMARKS_WHITE)
                        elif card == "YELLOW":
                            remarks = random.choice(REMARKS_YELLOW)
                        else:
                            remarks = random.choice(REMARKS_RED)
                        pts = points(dis, rp_main, rp_sec, pen_under, pen_other)
                        if dis in DEPTH:
                            secs = rp_main*2 + random.randint(-10, 10)
                            dive_time = "%d:%02d" % (secs // 60, secs % 60)

                ws.append([dis, last + " " + first, aid, gender, country,
                           ap_main, ap_sec,
                           rp_main, rp_sec,
                           pen_under, pen_other,
                           card, remarks, pts,
                           ot, wt, lane,
                           dive_time, None, None])
                # real exports link the diver name to the AIDA profile
                name_cell = ws.cell(row=ws.max_row, column=2)
                name_cell.hyperlink = PROFILE_URL + aid
                name_cell.font = Font(color="0563C1", underline="single")

                lane += 1
                if lane > n_lanes:
                    lane = 1
                    ot_min += 15 if dis != "STA" else 12
                    if ot_min >= 60:
                        ot_min -= 60
                        ot_hour += 1

    wb.save(outfile)
    print("Wrote", outfile)
    print("Competition days:", ", ".join(d.isoformat() for d in days))
    print("Athletes:", len(athletes))


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else "test_competition.xlsx")
